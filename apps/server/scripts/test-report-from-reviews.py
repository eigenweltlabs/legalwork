"""Regression checks for the single-call DD report builder."""
import hashlib
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest
from docx import Document

script = Path(__file__).resolve().parents[1] / "resources/core-opencode/skills/docx-edit/assets/report-from-reviews.py"
spec = importlib.util.spec_from_file_location("report_from_reviews", script)
helper = importlib.util.module_from_spec(spec)
spec.loader.exec_module(helper)


class ReportBuilderChecks(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.project = Path(self.directory.name)
        (self.project / "agreement.pdf").write_bytes(b"original source")
        source_hash = hashlib.sha256(b"original source").hexdigest()
        self.save("evidence.json", {"document": "agreement.pdf", "sourceHash": source_hash,
                                   "passages": [{"page": 2, "text": "A share transfer requires lender consent."}]})
        self.packet = {"distinctDocuments": 1, "openDocuments": 1, "openCells": 2,
                       "measuredFields": {"c_rows": 1, "c_open": 1},
                       "reviews": [{"unresolved": [{"sourceRef": "S00001", "column": "status", "status": "needs_review"}]}],
                       "sources": [{"ref": "S00001", "document": "agreement.pdf", "title": "Financing agreement",
                                    "sourceHash": source_hash, "evidenceFiles": ["evidence.json"]}]}
        self.draft = {"fields": {"executive_summary": "Obtain the lender's consent before completion.", "reviewed_files": 999,
                                 "unresolved_files": 0, "report_status": "Final", "c_rows": 99, "c_open": 0},
                      "findings": [{"title": "Lender consent", "severity": "High", "impact": "Consent is required for the proposed transfer.",
                                    "action": "Obtain written consent.", "owner": "Seller", "timing": "Before completion",
                                    "citations": [{"sourceRef": "S00001", "page": 2, "quote": "A share transfer requires lender consent."}]}]}
        template = Document()
        template.add_paragraph("{{executive_summary}}")
        template.add_paragraph("{{report_status}}: {{reviewed_files}} reviewed; {{unresolved_files}} open; {{c_rows}} / {{c_open}}.")
        for tokens in [["sev", "finding", "action"], ["rid", "risk_finding", "risk_source", "risk_action"],
                       ["aid", "action_timing", "action_owner", "action_deliverable"]]:
            table = template.add_table(rows=2, cols=len(tokens))
            for cell, token in zip(table.rows[1].cells, tokens):
                cell.text = "{{" + token + "}}"
        template.sections[0].footer.paragraphs[0].text = "{{report_status}}"
        template.save(self.project / "template.docx")

    def save(self, path, value):
        (self.project / path).write_text(json.dumps(value))

    def build(self):
        self.save("packet.json", self.packet)
        self.save("draft.json", self.draft)
        return helper.build(self.project, "packet.json", "draft.json", self.project / "template.docx", "reports/report.docx")

    def test_one_build_resolves_references_verifies_quotes_and_preserves_open_questions(self):
        result = self.build()
        self.assertEqual(result["verifiedCitations"], 1)
        self.assertEqual(result["openDocuments"], 1)
        doc = Document(self.project / "reports/report.docx")
        self.assertIn("Draft: 1 reviewed; 1 open; 1 / 1.", [p.text for p in doc.paragraphs])
        self.assertEqual(doc.sections[0].footer.paragraphs[0].text, "Draft")
        self.assertIn("Financing agreement (agreement.pdf), p. 2", doc.tables[1].rows[1].cells[2].text)
        unresolved = json.loads((self.project / "reports/report-support/unresolved.json").read_text())
        self.assertEqual(unresolved["openCells"], 2)

    def test_unsupported_quote_prevents_report_creation(self):
        self.draft["findings"][0]["citations"][0]["quote"] = "A transfer does not require consent."
        with self.assertRaisesRegex(ValueError, "source-support"):
            self.build()
        self.assertFalse((self.project / "reports/report.docx").exists())

    def test_changed_source_is_not_verified(self):
        (self.project / "agreement.pdf").write_bytes(b"changed source")
        with self.assertRaisesRegex(ValueError, "Source changed"):
            self.build()

    def test_aggregate_finding_cannot_hide_an_uncited_affected_source(self):
        second = {**self.packet["sources"][0], "ref": "S00002", "document": "second.pdf"}
        self.packet["sources"].append(second)
        self.draft["findings"][0]["affectedSources"] = ["S00001", "S00002"]
        with self.assertRaisesRegex(ValueError, "source-support"):
            self.build()

    def test_unknown_reference_and_output_boundary_are_rejected(self):
        self.draft["findings"][0]["citations"][0]["sourceRef"] = "invented"
        with self.assertRaisesRegex(ValueError, "source"):
            self.build()
        with self.assertRaisesRegex(ValueError, "inside the project"):
            helper.project_file(self.project, "../outside.json")

    def test_a_draft_can_have_no_material_findings_without_clearing_open_decisions(self):
        self.draft["findings"] = []
        result = self.build()
        self.assertEqual(result["verifiedCitations"], 0)
        self.assertEqual(result["openDocuments"], 1)

    def test_group_citations_expand_800_sources_without_model_generated_id_lists(self):
        base = self.packet["sources"][0]
        self.packet["sources"] = [{**base, "ref": "S" + str(i), "document": str(i) + ".pdf"} for i in range(800)]
        self.packet["reviews"] = [{"reviewId": "customers", "rows": [
            {"sourceRef": source["ref"], "answers": [{"column": "consent", "accepted": True, "value": "Yes"}]}
            for source in self.packet["sources"]]}]
        self.draft["findings"][0]["citations"] = []
        self.draft["findings"][0]["citationGroups"] = [{"reviewId": "customers", "column": "consent", "values": ["Yes"], "page": 2,
                                                       "quote": "A share transfer requires lender consent."}]
        plan, register, _ = helper.assemble(self.packet, self.draft)
        self.assertEqual(len(register["findings"][0]["citations"]), 800)
        self.assertEqual(len(register["findings"][0]["affected_files"]), 800)
        self.assertLess(len(plan["tables"][1]["rows"][0]["risk_source"]), 400)

    def test_native_group_never_expands_uncertain_answers(self):
        self.packet["reviews"] = [{"reviewId": "customers", "rows": [{"sourceRef": "S00001",
            "answers": [{"column": "consent", "accepted": False, "value": "Yes"}]}]}]
        self.draft["findings"][0]["citationGroups"] = [{"reviewId": "customers", "column": "consent", "values": ["Yes"], "page": 2,
                                                       "quote": "A share transfer requires lender consent."}]
        with self.assertRaisesRegex(ValueError, "no accepted sources"):
            helper.assemble(self.packet, self.draft)

    def test_partitioned_packet_keeps_all_source_rows_and_open_decisions(self):
        self.save("sources-1.json", self.packet["sources"])
        self.save("rows-1.json", [{"sourceRef": "S00001", "answers": []}])
        self.save("open-1.json", self.packet["reviews"][0]["unresolved"])
        self.packet["sources"] = []
        self.packet["sourceFiles"] = ["sources-1.json"]
        self.packet["reviews"][0] = {"rowFiles": ["rows-1.json"], "unresolvedFiles": ["open-1.json"]}
        result = self.build()
        self.assertEqual(result["verifiedCitations"], 1)
        self.assertEqual(result["openDocuments"], 1)
        unresolved = json.loads((self.project / "reports/report-support/unresolved.json").read_text())
        self.assertEqual(len(unresolved["items"]), 1)

    def test_a_draft_cannot_be_rebound_to_a_different_review_snapshot(self):
        self.packet["packetId"] = "current-snapshot"
        self.draft["packetId"] = "previous-snapshot"
        with self.assertRaisesRegex(ValueError, "pinned review packet"):
            self.build()


if __name__ == "__main__":
    unittest.main()
